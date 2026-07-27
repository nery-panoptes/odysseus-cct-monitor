import csv
import io
import re
import unicodedata
from pathlib import Path
from urllib.parse import parse_qs, urlencode, urlparse, urlunparse

from .cfg import sec
from .summarizer import extract_text_from_file
from .util import flatkey, normcnpj, root


COMPANY_SPECIFIC_TYPES = {"acordo", "termoAditivoAcordo"}
SHEET_NAME_FILTER_SOURCES = {"google_sheet_filter", "google_sheet_name_filter", "name_filter"}
SHEET_IGNORE_NAMES = {"", "-", "--", "bloq", "bloqueado", "bloqueada"}
STATE_HINTS = {
    "acre": "AC",
    "alagoas": "AL",
    "amazonas": "AM",
    "bahia": "BA",
    "ceara": "CE",
    "distrito federal": "DF",
    "goias": "GO",
    "maranhao": "MA",
    "mato grosso do sul": "MS",
    "mato grosso": "MT",
    "minas gerais": "MG",
    "para": "PA",
    "pernambuco": "PE",
    "piaui": "PI",
    "rio de janeiro": "RJ",
    "rio grande do norte": "RN",
    "rio grande do nort": "RN",
    "rio grande do sul": "RS",
    "santa catarina": "SC",
    "sao paulo": "SP",
    "tocantins": "TO",
}

CNPJ_TOKEN_RE = re.compile(
    r"(?<!\d)(?:\d{2}\.?\d{3}\.?\d{3}/?\d{4}-?\d{2}|\d{14})(?!\d)"
)


def _cell(value):
    if value is None:
        return ""

    if isinstance(value, float) and value.is_integer():
        return str(int(value))

    return str(value).strip()


def _header_key(value):
    return flatkey(value)


def _pick_column(headers, candidates):
    keys = {_header_key(item): item for item in headers}

    for candidate in candidates:
        found = keys.get(_header_key(candidate))

        if found:
            return found

    for header in headers:
        key = _header_key(header)

        for candidate in candidates:
            if _header_key(candidate) in key:
                return header

    return ""


def _strip_accents(value):
    text = unicodedata.normalize("NFD", str(value or ""))
    return "".join(ch for ch in text if unicodedata.category(ch) != "Mn")


def _words(value):
    text = _strip_accents(value).lower()
    return re.findall(r"[a-z0-9]+", text)


def _wordset(value):
    return set(_words(value))


def _tokens_from_name(value):
    value = str(value or "").strip()

    if not value:
        return set()

    keys = {_header_key(value)}
    pieces = re.split(r"\s+x\s+|(?<=[A-Za-z])x(?=[A-Za-z])|/|;|,|\||\+", value, flags=re.I)

    for piece in pieces:
        for subpiece in re.split(r"\s+-\s+", piece):
            key = _header_key(subpiece)

            if len(key) >= 3 and key not in {"df", "go", "to"}:
                keys.add(key)

    return {key for key in keys if key and key not in SHEET_IGNORE_NAMES}


def _candidate_virtual_aliases(row, linked_aliases=None):
    linked_aliases = linked_aliases or []
    parts = [
        row.get("nome") or "",
        row.get("apelido") or "",
        *(linked_aliases or []),
    ]
    text = " ".join(parts)
    words = _wordset(text)
    key = _header_key(text)
    aliases = set()

    for part in parts:
        aliases.update(_tokens_from_name(part))

    if "comercio" in words and ("varejista" in words or "varej" in key):
        aliases.add("sindcomvarejista")

    if "process" in key and "dados" in words:
        aliases.add("sindpd")

    if (
        ("informatica" in words or "inform" in key)
        and ("serv" in key or "emp" in key or "industria" in words)
    ):
        aliases.add("sindesei")

    if "asseio" in words or "conservacao" in words or "limpeza" in words:
        aliases.update({"seac", "seacons"})

    if "contabeis" in words or "contabilidade" in words:
        aliases.add("sescon")

    if "turismo" in words:
        aliases.update({"sindetur", "fretatur"})

    if "saude" in words or "hosp" in key or "hospital" in words:
        aliases.update({"sbh", "sindsaude"})

    if "metal" in key or "mecanica" in words or "metalurgico" in key:
        aliases.update({"simeb", "sindmetal"})

    if "construcao" in words or "contruc" in key:
        aliases.add("sinduscon")

    if "mobiliario" in words or "mob" in key or "marceneiros" in words:
        aliases.add("sindimam")

    if "farmacia" in words or "farmaceuticos" in words:
        aliases.add("sinfito")

    if "secretarias" in words or "secretarios" in words:
        aliases.add("sisdf")

    if "bombeiros" in words:
        aliases.add("sindbombeiros")

    if "cargas" in words or "sitratterdf" in aliases:
        aliases.update({"sitratterdf", "sittraterdf"})

    if "barb" in key or "cabel" in key or "beleza" in words or "bel" in key:
        aliases.add("sindbeleza")

    if "loterias" in words or "lotericas" in words:
        aliases.add("sindloterias")

    if "seguros" in words or "corret" in key:
        aliases.update({"sindsec", "sincor"})

    if "clubes" in words:
        aliases.add("sindclubes")

    if "publicitarios" in words or "publicit" in key:
        aliases.add("sinpublicitariosbrasilia")

    if "enfermeiros" in words:
        aliases.add("sindicatodosenfermeirosdodf")

    return aliases


def _configured_monitor_ufs(cfg):
    ufs = []

    for uf in sec(cfg, "monitor").get("ufs", []) or []:
        uf = str(uf or "").strip().upper()

        if len(uf) == 2 and uf not in ufs:
            ufs.append(uf)

    return set(ufs)


def _infer_candidate_uf(row):
    uf = str(row.get("uf_inferida") or "").strip().upper()

    if len(uf) == 2:
        return uf

    text = _strip_accents(" ".join([
        row.get("nome") or "",
        row.get("apelido") or "",
    ])).lower()
    compact = f" {text} "

    for hint, hint_uf in sorted(STATE_HINTS.items(), key=lambda item: len(item[0]), reverse=True):
        if hint in compact:
            return hint_uf

    for hint_uf in STATE_HINTS.values():
        if re.search(rf"(?<![a-z]){hint_uf.lower()}(?![a-z])", text):
            return hint_uf

    return ""


def filter_candidates_to_monitor_ufs(cfg, candidates):
    mcfg = sec(cfg, "monitor_source")

    if not bool(mcfg.get("respect_monitor_ufs", True)):
        return candidates

    allowed = _configured_monitor_ufs(cfg)

    if not allowed:
        return candidates

    filtered = []

    for row in candidates:
        uf = _infer_candidate_uf(row)

        if not uf or uf in allowed:
            filtered.append(row)

    return filtered


def _read_xlsx_rows(path, sheet_name=""):
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook[sheet_name] if sheet_name else workbook[workbook.sheetnames[0]]

    raw_rows = list(worksheet.iter_rows(values_only=True))

    header = []
    header_index = 0

    for idx, row in enumerate(raw_rows):
        values = [_cell(value) for value in row]

        if any(values):
            header = values
            header_index = idx
            break

    if not header:
        return []

    rows = []

    for row in raw_rows[header_index + 1:]:
        values = [_cell(value) for value in row]

        if not any(values):
            continue

        item = {}

        for idx, column in enumerate(header):
            if not column:
                continue

            item[column] = values[idx] if idx < len(values) else ""

        rows.append(item)

    return rows


def google_sheet_csv_url(url):
    url = str(url or "").strip()

    if not url:
        return ""

    parsed = urlparse(url)

    if "docs.google.com" not in parsed.netloc or "/spreadsheets/d/" not in parsed.path:
        return url

    if "/export" in parsed.path and "format=csv" in parsed.query:
        return url

    spreadsheet_id = parsed.path.split("/spreadsheets/d/", 1)[1].split("/", 1)[0]
    query = parse_qs(parsed.query)
    fragment_query = parse_qs(parsed.fragment)
    gid = (query.get("gid") or fragment_query.get("gid") or [""])[0]

    export_query = {"format": "csv"}

    if gid:
        export_query["gid"] = gid

    return urlunparse((
        "https",
        "docs.google.com",
        f"/spreadsheets/d/{spreadsheet_id}/export",
        "",
        urlencode(export_query),
        "",
    ))


def _read_csv_url_rows(url):
    try:
        import requests
    except ModuleNotFoundError as err:
        raise RuntimeError("A dependencia requests nao esta instalada.") from err

    csv_url = google_sheet_csv_url(url)
    response = requests.get(csv_url, timeout=60)

    if response.status_code != 200:
        raise RuntimeError(
            f"Planilha online retornou HTTP {response.status_code}. "
            "Confirme se o link esta publicado/compartilhado para leitura."
        )

    text = response.text

    if "<html" in text[:1000].lower() or "<!doctype html" in text[:1000].lower():
        raise RuntimeError(
            "O link da planilha online retornou HTML, nao CSV. "
            "Provavelmente a planilha esta privada ou o link nao e exportavel."
        )

    reader = csv.DictReader(io.StringIO(text))
    return [{key: _cell(value) for key, value in row.items()} for row in reader]


def _normalize_document(value):
    digits = normcnpj(value)

    if len(digits) == 13:
        digits = digits.zfill(14)

    if len(digits) == 14:
        return digits, "cnpj"

    if len(digits) == 11:
        return digits, "cpf"

    return "", ""


def _is_truthy_flag(value):
    text = _header_key(value)

    if not text:
        return True

    return text not in {"0", "n", "nao", "no", "false", "falso", "inativo", "inativa"}


def company_rows_from_table(rows, origem):
    if not rows:
        return []

    headers = list(rows[0].keys())
    id_col = _pick_column(headers, ["id", "codigo", "codigo cliente", "cod cliente"])
    name_col = _pick_column(headers, ["razao social", "razão social", "empresa", "cliente", "nome"])
    doc_col = _pick_column(headers, ["cnpj", "cpf cnpj", "cpf/cnpj", "documento", "cpf"])

    if not doc_col:
        raise RuntimeError("Nao encontrei coluna de CNPJ/CPF na base de empresas.")

    companies = []
    seen = set()

    for row in rows:
        document, document_type = _normalize_document(row.get(doc_col, ""))

        if not document or (document, origem) in seen:
            continue

        seen.add((document, origem))
        companies.append({
            "source_key": f"{origem}:{document}",
            "empresa_id": row.get(id_col, "") if id_col else "",
            "razao_social": row.get(name_col, "") if name_col else "",
            "documento": document,
            "tipo_documento": document_type,
            "origem": origem,
            "ativo": 1,
        })

    return companies


def load_office_company_rows(cfg):
    base = Path(cfg["base"])
    ocfg = sec(cfg, "office_companies")

    if not bool(ocfg.get("enabled", True)):
        return []

    rows = []
    source_file = str(ocfg.get("source_file", "") or "").strip()
    sheet_name = str(ocfg.get("sheet_name", "") or "").strip()

    if source_file:
        path = root(base, source_file)

        if path.exists():
            rows.extend(company_rows_from_table(_read_xlsx_rows(path, sheet_name), "empresas_xlsx"))
        else:
            raise RuntimeError(f"Base de empresas nao encontrada: {path}")

    google_url = str(ocfg.get("google_sheet_url", "") or "").strip()

    if google_url:
        rows.extend(company_rows_from_table(_read_csv_url_rows(google_url), "empresas_google_sheets"))

    return rows


def refresh_office_companies(cfg, db):
    rows = load_office_company_rows(cfg)

    if rows:
        origins = sorted({row["origem"] for row in rows})

        for origin in origins:
            db.replace_office_companies(
                [row for row in rows if row["origem"] == origin],
                origin,
            )

    return db.office_companies(document_type="cnpj")


def monitor_rows_from_table(rows):
    if not rows:
        return []

    headers = list(rows[0].keys())
    name_col = _pick_column(headers, ["sindicato", "sindicato nome", "nome", "razao social", "razão social"])
    cnpj_col = _pick_column(headers, ["cnpj", "cpf cnpj", "cpf/cnpj", "documento"])
    uf_col = _pick_column(headers, ["uf", "estado"])
    active_col = _pick_column(headers, ["monitorar", "ativo", "usar", "monitoramento"])

    if not cnpj_col:
        raise RuntimeError("Nao encontrei coluna de CNPJ na base online de sindicatos.")

    candidates = []
    seen = set()

    for idx, row in enumerate(rows, start=1):
        if active_col and not _is_truthy_flag(row.get(active_col, "")):
            continue

        document, document_type = _normalize_document(row.get(cnpj_col, ""))

        if document_type != "cnpj" or document in seen or document == "00000000000000":
            continue

        seen.add(document)
        candidates.append({
            "id": idx,
            "codigo": "",
            "nome": row.get(name_col, "") if name_col else "",
            "apelido": "",
            "cnpj": document,
            "uf_inferida": (row.get(uf_col, "") if uf_col else "").upper(),
            "monitorar_sugerido": "sim",
            "motivo_monitoramento": "base_online_google_sheets",
            "ocorrencias_atual": 0,
            "ocorrencias_2025": 0,
            "data_bases": "",
            "status_encontrados": "",
        })

    return candidates


def sheet_sindicato_names_from_table(rows, column_name=""):
    if not rows:
        return []

    headers = list(rows[0].keys())
    column = ""

    if column_name:
        requested = _header_key(column_name)

        for header in headers:
            if _header_key(header) == requested:
                column = header
                break

    if not column:
        column = _pick_column(headers, ["nome do sindicato", "nome sindicato", "sindicato"])

    if not column:
        raise RuntimeError("Nao encontrei a coluna NOME SINDICATO na planilha online.")

    names = []
    seen = set()

    for row in rows:
        value = str(row.get(column, "") or "").strip()
        key = _header_key(value)

        if key in SHEET_IGNORE_NAMES or key in seen:
            continue

        seen.add(key)
        names.append(value)

    return names


def _linked_aliases_by_candidate(db, candidates):
    cnpjs = {row.get("cnpj") for row in candidates if row.get("cnpj")}
    names = {row.get("nome") for row in candidates if row.get("nome")}
    aliases = {}

    if not cnpjs and not names:
        return aliases

    try:
        rows = db.con.execute("""
            select
                sindicato_nome_referencia,
                cadastro_cnpj_match,
                cadastro_nome_match
            from sindicatos_aliases
            where coalesce(cadastro_cnpj_match, '') != ''
               or coalesce(cadastro_nome_match, '') != ''
        """).fetchall()
    except Exception:
        return aliases

    for row in rows:
        ref = row["sindicato_nome_referencia"] or ""
        cnpj = row["cadastro_cnpj_match"] or ""
        name = row["cadastro_nome_match"] or ""

        if cnpj in cnpjs:
            aliases.setdefault(("cnpj", cnpj), []).append(ref)

        if name in names:
            aliases.setdefault(("nome", name), []).append(ref)

    return aliases


def _candidate_matches_tokens(row, allowed_tokens, linked_aliases=None):
    linked_aliases = linked_aliases or []
    fields = [
        row.get("nome") or "",
        row.get("apelido") or "",
        *(linked_aliases or []),
    ]
    candidate_keys = _candidate_virtual_aliases(row, linked_aliases=linked_aliases)
    candidate_text_key = _header_key(" ".join(fields))
    candidate_word_tokens = set()

    for field in fields:
        candidate_word_tokens.update(_words(field))

    for token in allowed_tokens:
        if not token:
            continue

        if token in candidate_keys:
            return True

        if len(token) >= 6 and (token in candidate_text_key or candidate_text_key in token):
            return True

        if 3 <= len(token) <= 5 and token in candidate_word_tokens:
            return True

    return False


def filter_monitor_candidates_by_sheet_names(db, candidates, sheet_names):
    allowed_tokens = set()

    for name in sheet_names:
        allowed_tokens.update(_tokens_from_name(name))

    aliases = _linked_aliases_by_candidate(db, candidates)
    filtered = []
    unmatched_names = []

    for row in candidates:
        linked_aliases = []
        linked_aliases.extend(aliases.get(("cnpj", row.get("cnpj") or ""), []))
        linked_aliases.extend(aliases.get(("nome", row.get("nome") or ""), []))

        if _candidate_matches_tokens(row, allowed_tokens, linked_aliases=linked_aliases):
            filtered.append(row)

    for name in sheet_names:
        tokens = _tokens_from_name(name)
        matched = False

        for row in filtered:
            linked_aliases = []
            linked_aliases.extend(aliases.get(("cnpj", row.get("cnpj") or ""), []))
            linked_aliases.extend(aliases.get(("nome", row.get("nome") or ""), []))

            if _candidate_matches_tokens(row, tokens, linked_aliases=linked_aliases):
                matched = True
                break

        if not matched:
            unmatched_names.append(name)

    return filtered, {
        "sheet_names": sheet_names,
        "allowed_tokens": sorted(allowed_tokens),
        "unmatched_names": unmatched_names,
    }


def monitor_name_filter_report(cfg, db, only_with_cnpj=True):
    mcfg = sec(cfg, "monitor_source")
    url = str(mcfg.get("google_sheet_url", "") or "").strip()
    filter_column = str(mcfg.get("filter_column", "") or "").strip()

    if not url:
        raise RuntimeError("google_sheet_url nao esta configurado em [monitor_source].")

    sheet_rows = _read_csv_url_rows(url)
    sheet_names = sheet_sindicato_names_from_table(sheet_rows, filter_column)
    raw_candidates = db.monitor_candidates(only_with_cnpj=only_with_cnpj)
    candidates = filter_candidates_to_monitor_ufs(cfg, raw_candidates)
    filtered, info = filter_monitor_candidates_by_sheet_names(db, candidates, sheet_names)

    return {
        "sheet_names": sheet_names,
        "total_sheet_names": len(sheet_names),
        "total_candidates_before": len(raw_candidates),
        "total_candidates_after_uf_scope": len(candidates),
        "total_candidates_after": len(filtered),
        "candidates": filtered,
        "unmatched_names": info["unmatched_names"],
        "csv_url": google_sheet_csv_url(url),
        "filter_column": filter_column or "NOME SINDICATO",
    }


def load_monitor_candidates(cfg, db, only_with_cnpj=True):
    mcfg = sec(cfg, "monitor_source")
    source = str(mcfg.get("source", "database") or "database").strip().lower()
    sheet_name = str(mcfg.get("sheet_name", "") or "").strip()
    filter_column = str(mcfg.get("filter_column", "") or "").strip()

    if source in SHEET_NAME_FILTER_SOURCES:
        url = str(mcfg.get("google_sheet_url", "") or "").strip()

        if not url:
            return db.monitor_candidates(only_with_cnpj=only_with_cnpj), "banco SQLite"

        sheet_rows = _read_csv_url_rows(url)
        sheet_names = sheet_sindicato_names_from_table(sheet_rows, filter_column)
        raw_candidates = db.monitor_candidates(only_with_cnpj=only_with_cnpj)
        candidates = filter_candidates_to_monitor_ufs(cfg, raw_candidates)
        filtered, info = filter_monitor_candidates_by_sheet_names(db, candidates, sheet_names)
        return filtered, (
            "Google Sheets filtrando nomes "
            f"({len(filtered)}/{len(raw_candidates)} candidatos; "
            f"{len(info['sheet_names'])} nomes unicos na planilha)"
        )

    if source in {"google_sheet", "google_sheets", "sheets"}:
        url = str(mcfg.get("google_sheet_url", "") or "").strip()

        if url:
            rows = monitor_rows_from_table(_read_csv_url_rows(url))
            return rows, "Google Sheets"

    if source in {"xlsx", "excel"}:
        file_name = str(mcfg.get("source_file", "") or "").strip()

        if file_name:
            rows = monitor_rows_from_table(_read_xlsx_rows(root(Path(cfg["base"]), file_name), sheet_name))
            return rows, "XLSX"

    return db.monitor_candidates(only_with_cnpj=only_with_cnpj), "banco SQLite"


def is_valid_cnpj(value):
    digits = normcnpj(value)

    if len(digits) != 14 or digits == digits[0] * 14:
        return False

    def check_digit(part, weights):
        total = sum(int(digit) * weight for digit, weight in zip(part, weights))
        mod = total % 11
        return "0" if mod < 2 else str(11 - mod)

    first = check_digit(digits[:12], [5, 4, 3, 2, 9, 8, 7, 6, 5, 4, 3, 2])
    second = check_digit(digits[:12] + first, [6, 5, 4, 3, 2, 9, 8, 7, 6, 5, 4, 3, 2])

    return digits[-2:] == first + second


def _compact_text(value):
    return re.sub(r"\s+", " ", str(value or "")).strip(" ,.;:-")


def _document_signatory_text(text):
    clean = str(text or "")
    normalized = _strip_accents(clean).lower()
    markers = []

    for pattern in (r"\bcelebram\b", r"\bcelebra\b", r"\bestipulando\b", r"\bclausula primeira\b"):
        match = re.search(pattern, normalized)

        if match:
            markers.append(match.start())

    if markers:
        return clean[:min(markers)]

    return clean[:6000]


def _party_name_before_cnpj(text, cnpj_start):
    prefix = str(text or "")[max(0, cnpj_start - 700):cnpj_start]
    separators = []

    for pattern in (
        r"\n\s*E\s*\n",
        r";\s*E\s+",
        r"\n\s*E\s+",
        r"\s+E\s*\n",
        r";",
        r"\n{2,}",
    ):
        separators.extend(match.end() for match in re.finditer(pattern, prefix, flags=re.I))

    if separators:
        prefix = prefix[max(separators):]

    for marker in (
        "Confira a autenticidade",
        "sistemas/mediador/.",
        "DATA DO PROTOCOLO",
    ):
        pos = _strip_accents(prefix).lower().rfind(_strip_accents(marker).lower())

        if pos >= 0:
            prefix = prefix[pos + len(marker):]

    prefix = re.sub(r"CNPJ\s*(?:n\.?|n)?\s*$", "", prefix, flags=re.I)
    return _compact_text(prefix)


def _looks_like_union_party(name):
    words = _wordset(name)
    key = _header_key(name)
    union_markers = {
        "sindicato",
        "sind",
        "federacao",
        "fed",
        "confederacao",
        "confed",
        "central",
    }

    if words & union_markers:
        return True

    return key.startswith(("sind", "fed", "confed"))


def extract_document_signatory_companies(text, sindicato_cnpj=""):
    signatory_text = _document_signatory_text(text)
    sindicato_cnpj = normcnpj(sindicato_cnpj)
    companies = []
    seen = set()

    for match in CNPJ_TOKEN_RE.finditer(signatory_text):
        digits = normcnpj(match.group(0))

        if len(digits) != 14 or not is_valid_cnpj(digits) or digits in seen:
            continue

        seen.add(digits)
        name = _party_name_before_cnpj(signatory_text, match.start())

        if digits == sindicato_cnpj:
            continue

        if _looks_like_union_party(name):
            continue

        companies.append({
            "razao_social": name,
            "documento": digits,
        })

    return companies


def extract_document_cnpjs(path, cfg=None):
    result = extract_text_from_file(path, cfg=cfg)
    text = result.get("text") or ""
    cnpjs = []

    for match in CNPJ_TOKEN_RE.finditer(text):
        digits = normcnpj(match.group(0))

        if len(digits) == 14 and is_valid_cnpj(digits) and digits not in cnpjs:
            cnpjs.append(digits)

    return {
        "ok": bool(result.get("ok")),
        "cnpjs": cnpjs,
        "text": text,
        "file_type": result.get("file_type") or "",
        "ocr_needed": bool(result.get("ocr_needed")),
        "warnings": result.get("warnings") or [],
    }


def company_specific_types(cfg):
    values = sec(cfg, "office_companies").get("company_specific_types", [])

    if not values:
        return set(COMPANY_SPECIFIC_TYPES)

    return {str(value or "").strip() for value in values if str(value or "").strip()}


def company_filter_enabled(cfg):
    return bool(sec(cfg, "office_companies").get("enabled", True))


def _company_lookup(companies):
    return {row.get("documento", ""): row for row in companies if row.get("documento")}


def evaluate_company_filter(cfg, item, file_path, office_companies):
    ocfg = sec(cfg, "office_companies")
    alert_unknown = bool(ocfg.get("alert_when_company_unknown", True))

    decision = {
        "status": "not_applicable",
        "should_alert": True,
        "document_cnpjs": [],
        "company_cnpjs": [],
        "company_parties": [],
        "matched_companies": [],
        "reason": "",
    }

    if not company_filter_enabled(cfg):
        decision["status"] = "disabled"
        return decision

    instrument_type = str(item.get("tipo_instrumento") or "").strip()

    if instrument_type not in company_specific_types(cfg):
        return decision

    if not file_path:
        decision.update({
            "status": "unknown",
            "should_alert": alert_unknown,
            "reason": "documento nao baixado; nao foi possivel extrair CNPJs",
        })
        return decision

    extracted = extract_document_cnpjs(file_path, cfg=cfg)
    document_cnpjs = extracted["cnpjs"]
    sindicato_cnpj = normcnpj(item.get("sindicato_cnpj") or "")
    signatory_companies = extract_document_signatory_companies(
        extracted.get("text") or "",
        sindicato_cnpj=sindicato_cnpj,
    )
    company_cnpjs = [company["documento"] for company in signatory_companies]

    if not company_cnpjs:
        company_cnpjs = [cnpj for cnpj in document_cnpjs if cnpj != sindicato_cnpj]

    lookup = _company_lookup(office_companies)
    matched = [lookup[cnpj] for cnpj in company_cnpjs if cnpj in lookup]

    decision["document_cnpjs"] = document_cnpjs
    decision["company_cnpjs"] = company_cnpjs
    decision["company_parties"] = signatory_companies
    decision["matched_companies"] = matched

    if matched:
        decision["status"] = "matched"
        decision["reason"] = "empresa do documento encontrada na base do escritorio"
        return decision

    if company_cnpjs:
        decision["status"] = "not_matched"
        decision["should_alert"] = False
        decision["reason"] = "empresa do documento nao pertence a base do escritorio"
        return decision

    decision.update({
        "status": "unknown",
        "should_alert": alert_unknown,
        "reason": "documento sem CNPJ de empresa identificado",
    })
    return decision


def apply_company_filter(db, cfg, item, instrumento_id, file_path, office_companies):
    decision = evaluate_company_filter(cfg, item, file_path, office_companies)

    item["empresa_filter_status"] = decision["status"]
    item["empresas_documento_cnpjs"] = decision["company_cnpjs"]
    item["empresas_documento_partes"] = decision.get("company_parties") or []
    item["empresas_escritorio_match"] = decision["matched_companies"]
    item["empresa_filter_reason"] = decision["reason"]

    if instrumento_id:
        db.set_company_filter(
            instrumento_id,
            decision["status"],
            decision["company_cnpjs"],
            decision["matched_companies"],
            decision["reason"],
        )

    return decision


def format_company_matches(companies):
    labels = []

    for company in companies or []:
        name = str(company.get("razao_social") or "").strip()
        document = str(company.get("documento") or "").strip()

        if name and document:
            labels.append(f"{name} ({document})")
        elif name:
            labels.append(name)
        elif document:
            labels.append(document)

    return "; ".join(labels)
